#!/usr/bin/env python3
"""
Compile the CoupledMD manuscript markdown into a single-column .docx with the
upgraded figures embedded above their captions.

Usage:
    python build_docx.py <manuscript.md> <version>      e.g. 1.0.0

Writes:
    versions/manuscript_v<version>.docx
    versions/manuscript_v<version>.md   (snapshot of the source markdown)

python-docx is used (not pandoc) for precise figure placement and caption
control. Unicode super/subscripts (Na⁺, S², χ₁, α5, µs, Å, °) pass through and
render in Word's body font.
"""

import re
import shutil
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

HERE = Path(__file__).resolve().parent
FIGDIR = HERE / "figures"
VERSIONS = HERE / "versions"
VERSIONS.mkdir(exist_ok=True)

# figure number -> (file, render width in inches)
FIGFILES = {
    1: (FIGDIR / "figure_1.png", 6.5),
    2: (FIGDIR / "figure_2.png", 3.3),
    3: (FIGDIR / "figure_3.png", 6.5),
    4: (FIGDIR / "figure4_gprotein_interface.png", 6.5),
    5: (FIGDIR / "figure5_partner_switching.png", 6.5),
}

INK = RGBColor(0x17, 0x1b, 0x22)


# ── inline markdown (**bold**, *italic*, `code`) ──────────────────────────────
def parse_inline(text):
    out, buf = [], ""
    bold = italic = mono = False
    i = 0
    while i < len(text):
        if text[i:i + 2] == "**":
            if buf:
                out.append((buf, bold, italic, mono)); buf = ""
            bold = not bold; i += 2; continue
        if text[i] == "*":
            if buf:
                out.append((buf, bold, italic, mono)); buf = ""
            italic = not italic; i += 1; continue
        if text[i] == "`":
            if buf:
                out.append((buf, bold, italic, mono)); buf = ""
            mono = not mono; i += 1; continue
        buf += text[i]; i += 1
    if buf:
        out.append((buf, bold, italic, mono))
    return out


def add_runs(paragraph, text):
    for seg, bold, italic, mono in parse_inline(text):
        r = paragraph.add_run(seg)
        r.bold = bold
        r.italic = italic
        if mono:
            r.font.name = "Consolas"
            r.font.size = Pt(9)
    return paragraph


# ── styles ────────────────────────────────────────────────────────────────────
def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = INK
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15
    for name, size in [("Heading 1", 13), ("Heading 2", 11), ("Title", 16)]:
        st = doc.styles[name]
        st.font.name = "Arial"
        st.font.size = Pt(size)
        st.font.color.rgb = INK
        st.font.bold = True


def add_figure(doc, num, caption_lines):
    path, width = FIGFILES[num]
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(path), width=Inches(width))
    # caption: first line bold lead ("Figure N. Title"), rest normal
    cap = doc.add_paragraph()
    cap.paragraph_format.space_after = Pt(10)
    cap.paragraph_format.left_indent = Inches(0.0)
    first = caption_lines[0]
    m = re.match(r"^(\*\*Figure \d+\..*?\*\*)(.*)$", first)
    if m:
        add_runs(cap, m.group(1))
        if m.group(2).strip():
            add_runs(cap, " " + m.group(2).strip())
    else:
        add_runs(cap, first)
    for line in caption_lines[1:]:
        add_runs(cap, " " + line.strip())
    for run in cap.runs:
        run.font.size = Pt(9)


# ── main build ────────────────────────────────────────────────────────────────
def build(md_path, version):
    md_path = Path(md_path)
    lines = md_path.read_text().splitlines()

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)
    setup_styles(doc)

    mode = "body"          # body | figures | references
    fig_buf = []           # accumulating caption lines for current figure
    fig_num = None

    def flush_figure():
        nonlocal fig_buf, fig_num
        if fig_num is not None and fig_buf:
            add_figure(doc, fig_num, fig_buf)
        fig_buf = []
        fig_num = None

    for raw in lines:
        line = raw.rstrip()
        stripped = line.strip()

        if stripped == "---":
            continue

        # headings
        if line.startswith("## "):
            title = line[3:].strip()
            low = title.lower()
            if low.startswith("figure caption"):
                mode = "figures"
                doc.add_heading("Figures", level=1)
                continue
            if low.startswith("reference"):
                flush_figure()
                mode = "references"
                doc.add_heading("References", level=1)
                continue
            flush_figure()
            mode = "body"
            doc.add_heading(title, level=1)
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            continue
        if line.startswith("# "):
            p = doc.add_paragraph(style="Title")
            add_runs(p, line[2:].strip())
            continue

        if not stripped:
            if mode != "figures":
                continue
            # blank line inside figures section separates nothing structural
            continue

        if mode == "figures":
            if stripped.startswith("**Figure"):
                flush_figure()
                fm = re.match(r"\*\*Figure (\d+)\.", stripped)
                fig_num = int(fm.group(1)) if fm else None
                fig_buf = [stripped]
            else:
                fig_buf.append(stripped)
            continue

        if mode == "references":
            m = re.match(r"^(\d+)\.\s+(.*)$", stripped)
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, stripped if not m else f"{m.group(1)}. {m.group(2)}")
            for r in p.runs:
                r.font.size = Pt(9)
            continue

        # ordinary body paragraph
        add_runs(doc.add_paragraph(), stripped)

    flush_figure()

    out_docx = VERSIONS / f"manuscript_v{version}.docx"
    doc.save(str(out_docx))
    shutil.copy(md_path, VERSIONS / f"manuscript_v{version}.md")
    print(f"  wrote {out_docx}")
    return out_docx


if __name__ == "__main__":
    md = sys.argv[1] if len(sys.argv) > 1 else str(HERE / "nar_manuscript.md")
    ver = sys.argv[2] if len(sys.argv) > 2 else "1.0.0"
    build(md, ver)
