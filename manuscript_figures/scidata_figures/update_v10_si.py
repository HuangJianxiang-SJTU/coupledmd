#!/usr/bin/env python3
"""Revise v10_si.docx to match the validated v9 final-207 data package."""
from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parent
DOCX = ROOT / "v10_si.docx"


REPLACEMENTS = {
    "Supplementary Data S1 serializes the final system identifier": (
        "Supplementary Data S1 serializes the final 207-system identifier, source PDB identifier, receptor name and mapped accession/gene when available, "
        "GPCR class, G-protein family and subtype label, authoritative ligand name and chemical identifier when present, selected-replica counts and durations, "
        "force-field/protocol label, lipid composition, structural-provenance label, bilayer/trajectory flags and source-availability flags. It does not serialize "
        "activation state, biological assembly, retained-chain lists or retrieval dates because those fields are not supported consistently by the authoritative sources."
    ),
    "Supplementary Data S10 inventories the current plotting": (
        "Supplementary Data S10 inventories the current plotting, package-build and API code, repository-relative paths, byte sizes, SHA-256 values, languages and "
        "versions available in the package-build environment. No authoritative record of the exact historical AMBER or GROMACS engine versions was located; therefore "
        "the supplement does not claim exact simulation-engine versions."
    ),
    "Supplementary Data S4 is a 416-row source-inventory evidence table": (
        "Supplementary Data S4 is a 414-row source-inventory evidence table with one aggregate production-trajectory source row and one topology source row per final "
        "system. It reports limited source locators, existence and byte size observed during audit, and the three-replica/500-ns release scope. S4 is not a frozen archive "
        "manifest and does not claim one row per deposited physical file, repository paths, per-file checksums or DOI deposition."
    ),
    "Supplementary Data S7 is the 624-row selected-replica release ledger": (
        "Supplementary Data S7 is the 621-row selected-replica release ledger and records three 500-ns selected replicas per system. The complete 621-file trajectory "
        "manifest, matched per-file topology identifiers, checksums and a self-contained file-level QC package remain pending final manifest freeze and deposition."
    ),
    "Supplementary Data S5 contains 2,168 detected-pocket rows": (
        "Supplementary Data S5 contains 2,149 detected-pocket rows from 205 systems. Gi_7YK6 and Gi_8YIC have valid API pocket records with n_pockets=0 and are represented "
        "by two explicit available_zero_pockets rows rather than by numerical-zero pocket measurements or unavailable records. Detected rows report pocket size, mean frequency, "
        "orthosteric annotation, stored location, reader-facing anatomical zone and mapped GPCRdb generic positions."
    ),
    "Supplementary Data S6 contains 5,824 tidy system-interface-metric rows": (
        "Supplementary Data S6 contains 5,796 tidy system-interface-metric rows (207 systems x seven adjacent TM-helix pairs x four metrics). The metrics are penetration "
        "and penetration_p90 (angstrom), open_fraction (unitless fraction) and occupancy (lipid heavy atoms per frame). Every row reports the three replica values, their mean, "
        "stored interval bounds, replica count, unit and schema version. The 95% intervals were obtained by resampling the three replica summaries 1,000 times under a fixed seed."
    ),
    "The production-trajectory acceptance criteria": (
        "The production-trajectory acceptance criteria, orthosteric-pocket-recovery validation and reduced-visualization audit are described in the main Technical Validation. "
        "Supplementary Data S7 is the 621-row selected-replica ledger. Supplementary Data S8 is a 207-system readiness/QC summary with pointers to the underlying evidence; it is "
        "not a self-contained full QC package. Supplementary Data S9 summarizes current portal/API file and content availability. A valid zero-pocket JSON record counts as available, "
        "and API file presence is not interpreted as deposited-archive coverage."
    ),
    "Supplementary Data S3 contains one row for each of the 105 unique field names": (
        "Supplementary Data S3 contains one row for each of the 105 unique field names actually serialized across S1-S10. Every row reports field, entity_level, type, unit, "
        "requiredness, nullable, allowed_vocabulary, definition, example, provenance_authority and schema_version. Hypothetical archive-manifest fields are not listed."
    ),
    "Supplementary Data S2 contains the 13 unresolved records": (
        "Supplementary Data S2 combines the 13 unresolved records and the two excluded records, Gq_7E9W and Gq_7DWC, in one 15-row release-boundary table. "
        "Gq_7DWC is a duplicate source identity of Gq_8DWC (both MRGPRX1, Q96LB2). There is no separate "
        "exclusions file. Neither unresolved nor excluded records contribute to final-cohort derived tables."
    ),
    "Supplementary Table S3 distinguishes the molecular-record roles": (
        "Supplementary Table S3 distinguishes the molecular-record roles relevant to the current server and future deposition. The current Supplementary Data package contains "
        "aggregate source evidence and derived CSV records; frozen per-file format metadata, pairings, checksums and recommended-reader details remain pending archive-manifest freeze."
    ),
    "The current OpenAPI schema describes the final-208 API access layer": (
        "The current OpenAPI schema describes the final-207 API access layer. Supplementary Data S9 reports ten current portal/API record types using explicit file/content "
        "availability semantics, including valid zero-pocket JSON records. These counts describe the current server and do not demonstrate a deposited archive or frozen file manifest."
    ),
    "Supplementary Figure S2 |": (
        "Supplementary Figure S2 | Definition and aggregation of transmembrane-gateway records. (A) Seven gateways formed by adjacent TM-helix pairs and the associated "
        "record-identity requirements. (B) A lipid heavy atom is classified as wedged when it lies within 5.0 Å of both helices and inside the central TM z-band; a frame is open "
        "when the maximum inward penetration, pmax, is at least 0.5 Å. (C) Frame-level measurements are summarized per replica and combined across three replicas to produce "
        "penetration, penetration-P90, open-fraction and occupancy records with stored 95% intervals."
    ),
    "Cite both the Scientific Data article": (
        "Cite the Scientific Data article and identify the CoupledMD server access date. Once a DOI-linked immutable archive is deposited, also cite its exact DOI and version."
    ),
    "The CoupledMD portal (https://www.coupledmd.cn) and API documentation": (
        "The CoupledMD portal (https://www.coupledmd.cn) and API documentation (https://www.coupledmd.cn/api/docs) provide the current final-207 access layer. "
        "The frozen manifest, archive checksums and DOI-linked deposition remain pending; no archive-completeness, DOI, licensing or direct-download claim is made."
    ),
    "Current final-208 records are available": (
        "Current final-207 records are available at https://www.coupledmd.cn (API documentation: /api/docs). "
        "The frozen manifest, checksums and DOI-linked archive remain pending deposition; no archive-completeness, licensing or direct-download claim is made."
    ),
}


TABLE0 = [
    ["Assertion", "Validated result"],
    ["Unique final system identifiers", "207"],
    ["Selected production replicas", "621"],
    ["Total selected production sampling", "310,500 ns"],
    ["Final systems covered by each current portal/API record type", "207"],
    ["Unresolved/excluded IDs in final derived tables", "0"],
    ["Gq_7E9W in final cohort/API records", "0"],
    ["Gq_7DWC in final cohort/API records", "0"],
    ["Pocket records", "2,149 detected rows + 2 valid zero-pocket rows"],
    ["Gateway records", "5,796 system/interface/metric rows"],
    ["Frozen archive, DOI and archive checksums", "Pending final manifest freeze and deposition"],
    ["Package schema and integrity assertions", "Passed"],
]


TABLE2 = [
    ["S3 field", "Type", "Requiredness", "Meaning"],
    ["field", "string", "required", "Exact column name serialized in S1-S10"],
    ["entity_level", "string", "required", "Actual record level at which the field is interpreted"],
    ["type", "string", "required", "Logical serialized type"],
    ["unit", "string", "required", "Measurement unit or explicit unitless/not-applicable label"],
    ["requiredness", "enum", "required", "required or conditional"],
    ["nullable", "boolean", "required", "Whether an empty CSV value is permitted"],
    ["allowed_vocabulary", "string", "conditional", "Pipe-delimited closed vocabulary when applicable"],
    ["definition", "string", "required", "Reader-facing semantic definition"],
    ["example", "string", "conditional", "Example non-null serialized value"],
    ["provenance_authority", "string", "required", "Authoritative source or computation"],
    ["schema_version", "string", "required", "Schema identifier v9-final207.1"],
]


TABLE3 = [
    ["Record role", "Format", "Required pairing/source", "Current interpretation"],
    ["Production-trajectory source evidence", "NetCDF or TRR", "Matched source topology", "Aggregate S4 evidence; not a deposited per-file manifest"],
    ["Topology source evidence", "PRMTOP/PARM7 or TPR", "Production trajectory", "Aggregate S4 evidence; not a deposited per-file manifest"],
    ["Reduced reference", "PDB", "Reduced XTC", "Current portal visualization"],
    ["Reduced trajectory", "XTC", "Reduced PDB", "Visualization and lightweight inspection"],
    ["Pocket/gateway records", "JSON/NPZ and CSV", "Per-system schema version", "Current derived annotations"],
    ["Readiness/QC summary", "CSV", "Evidence pointers", "Release evidence summary, not full per-file QC"],
]


TABLE4 = [
    ["File", "Content", "Granularity"],
    ["S1", "Included-system inventory; ligand fields populated only from authoritative metadata", "207 systems"],
    ["S2", "Combined release-boundary table: 13 unresolved and two excluded records", "15 records"],
    ["S3", "Dictionary for every field actually serialized in S1-S10", "105 unique fields"],
    ["S4", "Production/topology source-inventory evidence; not an archive manifest", "414 rows; two per final system"],
    ["S5", "Pocket summaries with GPCRdb mappings and explicit valid zero-pocket records", "2,149 pocket rows + 2 zero-pocket rows"],
    ["S6", "Tidy gateway means, intervals and three replica values for four metrics", "5,796 system/interface/metric rows"],
    ["S7", "Selected-replica release ledger; not frozen file-level QC", "621 system/replica rows"],
    ["S8", "Final-cohort readiness/QC summary with evidence pointers", "207 systems"],
    ["S9", "Current portal/API file and content availability; not archive coverage", "10 record types"],
    ["S10", "Current code files, environment files and package-build dependencies", "24 inventory items"],
]


def remove_paragraph(paragraph) -> None:
    element = paragraph._element
    element.getparent().remove(element)


def replace_paragraphs(document: Document) -> None:
    for paragraph in list(document.paragraphs):
        text = paragraph.text.strip()
        if text.startswith("Supplementary Figure S3"):
            remove_paragraph(paragraph)
            continue
        for prefix, replacement in REPLACEMENTS.items():
            if text.startswith(prefix):
                paragraph.text = replacement
                break


def set_table(table, rows: list[list[str]]) -> None:
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)
    while len(table.rows) < len(rows):
        table.add_row()
    for row, values in zip(table.rows, rows):
        if len(values) != len(table.columns):
            raise ValueError((values, len(table.columns)))
        for cell, value in zip(row.cells, values):
            cell.text = value


def configure_tables(document: Document) -> None:
    for table in document.tables:
        for row_number, row in enumerate(table.rows):
            properties = row._tr.get_or_add_trPr()
            if properties.find(qn("w:cantSplit")) is None:
                properties.append(OxmlElement("w:cantSplit"))
            if row_number == 0 and properties.find(qn("w:tblHeader")) is None:
                header = OxmlElement("w:tblHeader")
                header.set(qn("w:val"), "true")
                properties.append(header)


def configure_paragraphs(document: Document) -> None:
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text.startswith(("Supplementary Note", "Supplementary Table", "Supplementary Figure", "Supplementary Data files", "Supplementary reuse cautions", "Supplementary Data Availability")):
            paragraph.paragraph_format.keep_with_next = True


def validate_text(document: Document) -> None:
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    forbidden = [
        "A separate final-exclusions table", "exact engine and dependency versions used",
        "activation-state assignment", "Each deposited pocket record", "Supplementary Figure S3",
    ]
    for phrase in forbidden:
        assert phrase not in text, phrase
    required = [
        "2,149 detected-pocket rows", "5,796 tidy system-interface-metric rows",
        "There is no separate exclusions file", "not a self-contained full QC package",
        "does not claim exact simulation-engine versions",
    ]
    for phrase in required:
        assert phrase in text, phrase


def main() -> None:
    document = Document(DOCX)
    replace_paragraphs(document)
    assert len(document.tables) == 5
    set_table(document.tables[0], TABLE0)
    set_table(document.tables[2], TABLE2)
    set_table(document.tables[3], TABLE3)
    set_table(document.tables[4], TABLE4)
    configure_tables(document)
    configure_paragraphs(document)
    document.core_properties.modified = datetime.now(timezone.utc)
    validate_text(document)
    document.save(DOCX)
    print(f"Updated {DOCX}")


if __name__ == "__main__":
    main()
